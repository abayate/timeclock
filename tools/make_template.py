"""Builds the starter invoice template (app/template/invoice-template.docx).

Placeholders in {curly braces} are filled in by the app. The table row
wrapped in {#shifts} ... {/shifts} repeats once per shift.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else (
    Path(__file__).resolve().parent.parent / "app" / "template" / "invoice-template.docx")
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def para(doc_or_cell, text="", bold=False, size=11, color=None, align=None, space_after=2):
    p = doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    para(doc, "INVOICE", bold=True, size=26, color=ACCENT, space_after=10)

    # From / invoice details side by side
    top = doc.add_table(rows=1, cols=2)
    left, right = top.rows[0].cells
    left._tc.remove(left.paragraphs[0]._p)
    right._tc.remove(right.paragraphs[0]._p)
    para(left, "FROM", bold=True, size=9, color=MUTED)
    para(left, "{my_name}", bold=True)
    para(left, "{my_email}")
    para(left, "")
    para(left, "BILL TO", bold=True, size=9, color=MUTED)
    para(left, "{bill_to}")
    for label, value in [
        ("Invoice #", "{invoice_number}"),
        ("Invoice date", "{invoice_date}"),
        ("Period", "{period}"),
    ]:
        p = para(right, align=WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(f"{label}: ")
        r.font.color.rgb = MUTED
        p.add_run(value).bold = True

    para(doc, "", space_after=8)

    # Hours table: header + one looping row
    headers = ["Date", "Clock in", "Clock out", "Hours"]
    fields = ["{#shifts}{date}", "{in}", "{out}", "{hours}{/shifts}"]
    t = doc.add_table(rows=2, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F4E79")
    for i, f in enumerate(fields):
        t.rows[1].cells[i].paragraphs[0].add_run(f)
    for row in t.rows:
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    para(doc, "", space_after=8)

    # Totals
    lines = [
        ("Total hours", "{total_hours}"),
        ("Hourly rate", "${rate}"),
        ("Subtotal", "${subtotal}"),
        ("Wire fee", "${wire_fee}"),
        ("Amount due", "${total}"),
    ]
    totals = doc.add_table(rows=len(lines), cols=2)
    totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for row, (label, value) in zip(totals.rows, lines):
        row.cells[0].paragraphs[0].add_run(label).font.color.rgb = MUTED
        p = row.cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(value)
        r.bold = True
        if label == "Amount due":
            r.font.size = Pt(14)
            r.font.color.rgb = ACCENT

    para(doc, "", space_after=12)
    para(doc, "{notes}", color=MUTED, size=10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
